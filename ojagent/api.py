"""FastAPI backend for the UI: upload a resume, start a triage run, follow its progress.

    uvicorn ojagent.api:app --port 8000

With the UI built (`npm run build` in ui/), the same server also serves it at /, which is
how the Docker image runs.

A run scrapes the job site the user picked (onlinejobs.ph, JobStreet or both) with a 5 s
delay between requests, so it takes minutes. It runs on a background thread; the UI
polls GET /api/runs/{id}. One run at a time keeps the crawl polite.
"""

from __future__ import annotations

import functools
import io
import json
import re
import threading
import time
import uuid
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Callable, Literal

from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field

from .config import ROOT, load_settings
from .llm import LLM, build_llm
from .prefilter import JevPrefilter
from .profile_store import JobSnapshot, Mark, Profile, ProfileStore, learned_summary, seed_from_markdown
from . import sources
from .summary import SummaryStore, source_hash, summarize
from .scraper import JobSource
from .triage import render, to_json, triage

MAX_RESUME_BYTES = 5 * 1024 * 1024
STAGE_LABELS = {
    "screen": "Jev is screening postings",
    "fetch": "Reading postings (5 s wait between requests)",
    "score": "Jev is scoring postings",
    "summarize": "Writing summaries",
    "done": "Done",
}


def extract_resume_text(filename: str, data: bytes) -> str:
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        from pypdf import PdfReader

        text = "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(data)).pages)
    elif ext == ".docx":
        import docx

        doc = docx.Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        parts += [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
        text = "\n".join(parts)
    elif ext in (".txt", ".md"):
        text = data.decode("utf-8", "replace")
    else:
        raise ValueError("Upload a PDF, DOCX, TXT or MD resume.")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if len(text) < 200:
        raise ValueError("Couldn't read enough text from that file. Is it a scanned image?")
    return text


class MarkRequest(BaseModel):
    status: Literal["applied", "dismissed"] | None  # None clears the mark
    job: JobSnapshot | None = None
    note: str = Field(default="", max_length=300)


class RunRequest(BaseModel):
    keyword: str = Field(default="python", min_length=1, max_length=60)
    postings: int = Field(default=30, ge=5, le=90)  # newest postings to scan, capped per source
    source: Literal["onlinejobs", "jobstreet", "all"] = "onlinejobs"


@dataclass
class Run:
    keyword: str
    postings: int
    source: str = "onlinejobs"
    id: str = field(default_factory=lambda: uuid.uuid4().hex[:8])
    status: str = "running"  # running | done | error
    stage: str = "search"
    done: int = 0
    total: int = 0
    started: float = field(default_factory=time.time)
    finished: float | None = None
    error: str | None = None
    result: dict | None = None

    def stage_label(self) -> str:
        if self.stage == "search":
            return f"Searching {sources.SOURCES[self.source][0]}"
        return STAGE_LABELS.get(self.stage, self.stage)

    def view(self) -> dict:
        return {
            "id": self.id,
            "keyword": self.keyword,
            "postings": self.postings,
            "source": self.source,
            "status": self.status,
            "stage": self.stage,
            "stage_label": self.stage_label(),
            "done": self.done,
            "total": self.total,
            "elapsed": round((self.finished or time.time()) - self.started, 1),
            "error": self.error,
            "result": self.result,
        }


def _default_jev() -> JevPrefilter | None:
    key = load_settings().typesafe_api_key
    return JevPrefilter(key) if key else None


@functools.cache  # one client for the app, so it remembers which free model last worked
def _default_llm() -> LLM | None:
    try:
        s = load_settings()
    except SystemExit:  # no LLM key: Jev-only, summaries come from the posting text
        return None
    return build_llm(s.model, s.api_key, s.base_url, s.fallback_models)


def create_app(
    make_source: Callable[[str], JobSource] = sources.make_source,
    make_jev: Callable[[], JevPrefilter | None] = _default_jev,
    make_llm: Callable[[], LLM | None] = _default_llm,
    profile_path: Path = ROOT / "profile.md",
    data_dir: Path = ROOT / "data",
    out_dir: Path = ROOT / "out",
    ui_dir: Path = ROOT / "ui" / "dist",
) -> FastAPI:
    app = FastAPI(
        title="ojagent",
        description="Job triage for onlinejobs.ph and JobStreet Philippines",
    )
    store = ProfileStore(data_dir, profile_path)
    summaries = SummaryStore(data_dir)
    summary_lock = threading.Lock()  # one summary write at a time: after a search or from the page

    def write_summary(llm: LLM) -> None:
        with summary_lock:
            summaries.save(summarize(llm, store.composed()))

    def summary_after_search(llm: LLM) -> None:
        """Bring the Profile page bullets up to date with the marks made since the last summary."""
        summary = summaries.load()
        if summary and summary.get("source_hash") == source_hash(store.composed()):
            return  # nothing changed, so a new summary would say the same
        try:
            write_summary(llm)
        except Exception:
            pass  # best effort: the page shows the old bullets and can retry
    runs: dict[str, Run] = {}
    lock = threading.Lock()

    def profile_view() -> dict:
        profile, marks = store.load(), store.marks()
        composed = store.composed()
        summary = summaries.load()
        return {
            "profile": profile.model_dump(),
            "learned": learned_summary(marks),
            "composed": composed,
            "summary": summary,
            # The profile, marks or reasons changed since the bullets were written.
            "summary_stale": summary is None or summary.get("source_hash") != source_hash(composed),
        }

    def execute(run: Run) -> None:
        def progress(stage: str, done: int, total: int) -> None:
            run.stage, run.done, run.total = stage, done, total

        try:
            jev, llm = make_jev(), make_llm()
            if jev is None and llm is None:
                raise RuntimeError("Set TYPESAFE_API_KEY (and optionally an LLM key) in .env")
            profile = store.composed()
            label = sources.SOURCES[run.source][0]
            result = triage(make_source(run.source), jev, profile, llm=llm, limit=run.postings,
                            keyword=run.keyword, progress=progress)
            data = to_json(result, profile)
            data["keyword"] = run.keyword
            data["source"] = run.source
            out_dir.mkdir(parents=True, exist_ok=True)
            stem = out_dir / f"triage-{datetime.now():%Y%m%d-%H%M%S}-{'jev' if jev else 'no-jev'}"
            stem.with_suffix(".md").write_text(render(result, profile, source_label=label), encoding="utf-8")
            text = json.dumps(data, ensure_ascii=False, indent=1)
            stem.with_suffix(".json").write_text(text, encoding="utf-8")
            (out_dir / "latest.json").write_text(text, encoding="utf-8")
            run.result, run.status = data, "done"
            if llm is not None:
                threading.Thread(target=summary_after_search, args=(llm,), daemon=True).start()
        except Exception as e:
            run.error, run.status = f"{type(e).__name__}: {e}", "error"
        finally:
            run.finished = time.time()

    @app.get("/api/profile")
    def get_profile() -> dict:
        return profile_view()

    @app.put("/api/profile")
    def put_profile(profile: Profile) -> dict:
        store.save(profile)
        return profile_view()

    @app.post("/api/profile/summary")
    def refresh_summary() -> dict:
        llm = make_llm()
        if llm is None:
            raise HTTPException(503, "Set an LLM key in .env to get a profile summary.")
        try:
            write_summary(llm)
        except Exception as e:  # bad JSON or a provider hiccup: keep the old bullets
            raise HTTPException(502, f"Couldn't write the summary: {e}") from e
        return profile_view()

    @app.get("/api/resume")
    def get_resume() -> dict:
        return {"resume": store.load().resume_meta}

    @app.post("/api/resume")
    async def upload_resume(file: UploadFile = File(...)) -> dict:
        data = await file.read(MAX_RESUME_BYTES + 1)
        if len(data) > MAX_RESUME_BYTES:
            raise HTTPException(413, "Resume must be under 5 MB.")
        try:
            text = extract_resume_text(file.filename or "", data)
        except ValueError as e:
            raise HTTPException(422, str(e))
        meta = {
            "filename": Path(file.filename or "resume").name,
            "uploaded_at": datetime.now().isoformat(timespec="seconds"),
            "chars": len(text),
        }
        profile = store.load()
        profile.resume_text, profile.resume_meta = text[:30_000], meta
        store.save(profile)
        return {"resume": meta, **profile_view()}

    @app.delete("/api/resume")
    def delete_resume() -> dict:
        """Go back to the resume text in profile.md."""
        profile = store.load()
        md = profile_path.read_text(encoding="utf-8") if profile_path.exists() else ""
        profile.resume_text, profile.resume_meta = seed_from_markdown(md).resume_text, None
        store.save(profile)
        return {"resume": None, **profile_view()}

    def marks_view(marks: dict[str, Mark]) -> dict:
        return {
            "marks": {k: m.status for k, m in marks.items()},
            "jobs": {k: m.job.model_dump() for k, m in marks.items()},
            "notes": {k: m.note for k, m in marks.items() if m.note},
        }

    @app.get("/api/marks")
    def get_marks() -> dict:
        return marks_view(store.marks())

    @app.put("/api/marks/{job_id}")
    def put_mark(job_id: str, req: MarkRequest) -> dict:
        if req.status is None:
            marks = store.set_mark(job_id, None)
        else:
            if req.job is None:
                raise HTTPException(422, "Send the job with an applied or dismissed mark.")
            marks = store.set_mark(job_id, Mark(status=req.status, job=req.job, note=req.note.strip()))
        return marks_view(marks)

    @app.delete("/api/marks")
    def clear_marks() -> dict:
        store.clear_marks()
        return {"marks": {}, "jobs": {}, "notes": {}}

    @app.post("/api/runs", status_code=202)
    def start_run(req: RunRequest) -> dict:
        cap = sources.SOURCES[req.source][2]
        if req.postings > cap:
            raise HTTPException(422, f"{sources.SOURCES[req.source][0]} allows at most {cap} postings per run.")
        with lock:
            if any(r.status == "running" for r in runs.values()):
                raise HTTPException(409, "A search is already running.")
            run = Run(keyword=req.keyword.strip(), postings=req.postings, source=req.source)
            runs[run.id] = run
        threading.Thread(target=execute, args=(run,), daemon=True).start()
        return run.view()

    @app.get("/api/runs/active")
    def active_run() -> dict:
        running = [r for r in runs.values() if r.status == "running"]
        return {"run": running[0].view() if running else None}

    @app.get("/api/runs/{run_id}")
    def get_run(run_id: str) -> dict:
        run = runs.get(run_id)
        if run is None:
            raise HTTPException(404, "Unknown run.")
        return run.view()

    @app.get("/api/latest")
    def latest() -> dict:
        path = out_dir / "latest.json"
        if not path.exists():
            raise HTTPException(404, "No results yet. Start a search.")
        return json.loads(path.read_text(encoding="utf-8"))

    # Mounted last so /api routes win. The UI routes by hash, so serving files is enough.
    if (ui_dir / "index.html").exists():
        app.mount("/", StaticFiles(directory=ui_dir, html=True), name="ui")

    return app


app = create_app()
